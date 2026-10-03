"""Full pipeline through /api/resume/upload with PDF, DOCX, and bad inputs (live LLM + job APIs)."""
import io, os
from docx import Document
from app.main import app

c = app.test_client()

def post(name, data):
    r = c.post("/api/resume/upload", data={"resume": (io.BytesIO(data), name)}, content_type="multipart/form-data")
    j = r.get_json()
    if r.status_code == 200:
        cp = j["candidate_profile"]
        return f"200 years={cp['total_experience_years']} level={cp['career_level']} recs={len(j['recommendations'])} queries={len(j['search_queries'])}"
    return f"{r.status_code} {j.get('error','')[:160]}"

pdf = open("C:/Users/dhair/Downloads/SivaSaiAkhilesh_MTS.pdf", "rb").read()
print("real PDF        ->", post("siva.pdf", pdf), flush=True)

d = Document()
for line in ["Priya Sharma  priya@example.com", "Education: B.Tech Computer Science 2026",
             "Skills: Python, Java, SQL, React, Flask",
             "Projects", "Library Management System - Flask and MySQL CRUD app.", "Weather Dashboard - React app using OpenWeather API.",
             "Experience", "Software Engineering Intern, Zoho   Jan 2026 - Jun 2026", "Built REST endpoints in Java Spring Boot and wrote unit tests."]:
    d.add_paragraph(line)
buf = io.BytesIO(); d.save(buf)
print("DOCX fresher    ->", post("priya.docx", buf.getvalue()), flush=True)

print("empty PDF bytes ->", post("bad.pdf", b""), flush=True)
print("corrupt PDF     ->", post("bad2.pdf", b"%PDF-1.4 not really a pdf"), flush=True)
print("txt file        ->", post("a.txt", b"hello"), flush=True)
d2 = Document(); d2.add_paragraph("hello"); b2 = io.BytesIO(); d2.save(b2)
print("tiny DOCX       ->", post("tiny.docx", b2.getvalue()), flush=True)
print("DONE", flush=True)
